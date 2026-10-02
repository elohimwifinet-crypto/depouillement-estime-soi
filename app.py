# -*- coding: utf-8 -*-
import os
import io
from datetime import datetime

from flask import (
    Flask, render_template, request, redirect, url_for, flash, send_file, abort
)
from flask_login import (
    LoginManager, UserMixin, login_user, logout_user, login_required,
    current_user
)
from werkzeug.security import generate_password_hash, check_password_hash

import db
from items import (
    SECTION_II, SECTION_III, SECTION_IV, SECTION_V, SECTION_B,
    AGE_CHOICES, SEXE_CHOICES, SITUATION_FAMILIALE_CHOICES,
    DATA_COLUMNS, build_dictionnaire_rows
)

app = Flask(__name__)
app.config["SECRET_KEY"] = os.environ.get("SECRET_KEY", "change-moi-en-production")

login_manager = LoginManager()
login_manager.init_app(app)
login_manager.login_view = "login"


class User(UserMixin):
    def __init__(self, row):
        self.id = row["id"]
        self.username = row["username"]
        self.role = row["role"]


@login_manager.user_loader
def load_user(user_id):
    conn = db.get_db()
    row = conn.execute("SELECT * FROM users WHERE id = ?", (user_id,)).fetchone()
    return User(row) if row else None


def admin_required(fn):
    from functools import wraps

    @wraps(fn)
    def wrapper(*args, **kwargs):
        if not current_user.is_authenticated or current_user.role != "admin":
            abort(403)
        return fn(*args, **kwargs)
    return wrapper


# ---------- Auth ----------

@app.route("/login", methods=["GET", "POST"])
def login():
    if request.method == "POST":
        username = request.form["username"].strip()
        password = request.form["password"]
        conn = db.get_db()
        row = conn.execute("SELECT * FROM users WHERE username = ?", (username,)).fetchone()
        if row and check_password_hash(row["password_hash"], password):
            login_user(User(row))
            return redirect(url_for("dashboard"))
        flash("Identifiants incorrects.", "error")
    return render_template("login.html")


@app.route("/logout")
@login_required
def logout():
    logout_user()
    return redirect(url_for("login"))


# ---------- Dashboard / liste ----------

@app.route("/")
@login_required
def dashboard():
    import json as _json
    from collections import defaultdict
    conn = db.get_db()

    raw = conn.execute(
        "SELECT id, date_soumission, nom, prenom, age, sexe, classe, etablissement, "
        "section2_json, sei_json "
        "FROM reponses ORDER BY LOWER(TRIM(nom)) ASC, LOWER(TRIM(prenom)) ASC"
    ).fetchall()

    students = []
    for r in raw:
        s2 = _json.loads(r["section2_json"]) if r["section2_json"] else []
        sei = _json.loads(r["sei_json"]) if r["sei_json"] else []

        pp2 = sum(1 for a in s2 if a.get("polarite") == "pos" and str(a.get("reponse","")).lower() == "oui")
        pn2 = sum(1 for a in s2 if a.get("polarite") == "neg" and str(a.get("reponse","")).lower() == "non")
        score2 = pp2 + pn2

        pp_sei = sum(1 for a in sei if a.get("polarite") == "vrai_pos" and str(a.get("reponse","")).lower() == "vrai")
        pn_sei = sum(1 for a in sei if a.get("polarite") == "vrai_neg" and str(a.get("reponse","")).lower() == "faux")
        score_sei = pp_sei + pn_sei

        total = score2 + score_sei

        if score2 <= 13: niv2 = "\u0054r\u00e8s faible"; col2 = "#c0392b"
        elif score2 <= 20: niv2 = "Faible"; col2 = "#e67e22"
        elif score2 <= 27: niv2 = "Moyen"; col2 = "#f1c40f"
        elif score2 <= 33: niv2 = "\u00c9lev\u00e9"; col2 = "#2ecc71"
        else: niv2 = "\u0054r\u00e8s \u00e9lev\u00e9"; col2 = "#27ae60"

        if score_sei <= 5: niv_sei = "\u0054r\u00e8s faible"; col_sei = "#c0392b"
        elif score_sei <= 8: niv_sei = "Faible"; col_sei = "#e67e22"
        elif score_sei <= 11: niv_sei = "Moyen"; col_sei = "#f1c40f"
        elif score_sei <= 13: niv_sei = "\u00c9lev\u00e9"; col_sei = "#2ecc71"
        else: niv_sei = "\u0054r\u00e8s \u00e9lev\u00e9"; col_sei = "#27ae60"

        students.append({
            "id": r["id"], "date": r["date_soumission"],
            "nom": r["nom"], "prenom": r["prenom"],
            "age": r["age"], "sexe": r["sexe"],
            "classe": r["classe"], "etablissement": r["etablissement"],
            "created_at": r["date_soumission"],
            "pp2": pp2, "pn2": pn2, "score2": score2, "niv2": niv2, "col2": col2,
            "pp_sei": pp_sei, "pn_sei": pn_sei, "score_sei": score_sei,
            "niv_sei": niv_sei, "col_sei": col_sei, "total": total,
        })

    groupes = defaultdict(list)
    for st in students:
        cle = (st["nom"].strip().lower(), st["prenom"].strip().lower())
        groupes[cle].append(st)
    doublons = {k: v for k, v in groupes.items() if len(v) > 1}
    nb_doublons = sum(len(v) for v in doublons.values())

    entretiens = conn.execute(
        "SELECT id, date_entretien, nom_eleve FROM entretiens "
        "ORDER BY LOWER(TRIM(nom_eleve)) ASC, id DESC"
    ).fetchall()

    return render_template(
        "dashboard.html",
        students=students,
        entretiens=entretiens,
        nb_entretiens=len(entretiens),
        nb_doublons=nb_doublons,
        doublons=doublons,
    )


# ---------- Fiche élève : création ----------

def _collect_form_answers():
    """Récupère toutes les réponses du formulaire dans un dict {code: valeur}."""
    answers = {}
    for code, _ in SECTION_II:
        answers[code] = request.form.get(code)
    for code, _, _, _ in SECTION_III:
        answers[code] = request.form.get(code)
    for code, _, _, _ in SECTION_IV:
        answers[code] = request.form.get(code)
    for code, _ in SECTION_V:
        answers[code] = request.form.get(code)
    for code, _ in SECTION_B:
        answers[code] = request.form.get(code)
    return answers


@app.route("/students/new", methods=["GET", "POST"])
@login_required
def student_new():
    if request.method == "POST":
        import json as _json

        # Section II : 40 items oui/non
        s2 = []
        total_oui_ii = 0
        for i, (code, libelle) in enumerate(SECTION_II, start=1):
            rep = (request.form.get(code) or "").lower()
            s2.append({"item": "item_" + str(i), "reponse": rep})
            if rep == "oui":
                total_oui_ii += 1

        # Section III
        iii_1 = request.form.get("III_1") or ""
        iii_2 = request.form.get("III_2") or ""
        iii_3 = request.form.get("III_3") or ""

        # Section IV
        s4 = {}
        for i in range(1, 13):
            s4["q" + str(i)] = request.form.get("IV_" + str(i)) or ""

        # Section V
        s5 = []
        for i in range(1, 6):
            s5.append({"item": "v" + str(i), "reponse": request.form.get("V_" + str(i)) or ""})

        # Section B (SEI) avec polarite Coopersmith
        from items import SEI_POLARITES
        sb = []
        total_vrai_b = 0
        for i in range(1, 16):
            code = "B_" + str(i)
            rep = (request.form.get(code) or "").lower()
            pol = SEI_POLARITES.get(code, "vrai_pos")
            sb.append({"item": "b" + str(i), "polarite": pol, "reponse": rep})
            if pol == "vrai_pos" and rep == "vrai":
                total_vrai_b += 1
            elif pol == "vrai_neg" and rep == "faux":
                total_vrai_b += 1

        conn = db.get_db()
        conn.execute(
            """INSERT INTO reponses
            (date_soumission, nom, prenom, age, sexe, classe, etablissement, situation_familiale,
             section2_json, section3_niveau, section3_q2, section3_q3,
             section4_json, section5_json, sei_json, score_estime_soi, score_max)
            VALUES (datetime('now'), ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)""",
            (
                request.form.get("nom"),
                request.form.get("prenom"),
                request.form.get("age"),
                request.form.get("sexe"),
                request.form.get("classe"),
                request.form.get("etablissement"),
                request.form.get("situation_familiale"),
                _json.dumps(s2, ensure_ascii=False),
                iii_1, iii_2, iii_3,
                _json.dumps(s4, ensure_ascii=False),
                _json.dumps(s5, ensure_ascii=False),
                _json.dumps(sb, ensure_ascii=False),
                total_oui_ii,
                55,
            ),
        )
        conn.commit()
        flash("Eleve enregistre.", "success")
        return redirect(url_for("dashboard"))
    return render_template(
        "student_form.html",
        student=None,
        section_ii=SECTION_II, section_iii=SECTION_III, section_iv=SECTION_IV,
        section_v=SECTION_V, section_b=SECTION_B,
        age_choices=AGE_CHOICES, sexe_choices=SEXE_CHOICES,
        situation_choices=SITUATION_FAMILIALE_CHOICES,
    )



@app.route("/students/<int:student_id>")
@login_required
def student_view(student_id):
    import json as _json
    conn = db.get_db()
    row = conn.execute("SELECT * FROM reponses WHERE id = ?", (student_id,)).fetchone()
    if not row:
        abort(404)

    # Reconstruire un dict compatible avec le template student_view.html
    student = dict(row)

    # Section II : liste de 40 items oui/non
    s2 = _json.loads(row["section2_json"]) if row["section2_json"] else []
    for i, item in enumerate(s2, start=1):
        student["II_" + str(i)] = item.get("reponse", "")

    # Section III
    student["III_1"] = row["section3_niveau"] or ""
    student["III_2"] = row["section3_q2"] or ""
    student["III_3"] = row["section3_q3"] or ""

    # Section IV : dict {q1:..., q2:..., ...}
    s4 = _json.loads(row["section4_json"]) if row["section4_json"] else {}
    for i in range(1, 13):
        student["IV_" + str(i)] = s4.get("q" + str(i), "")

    # Section V : liste
    s5 = _json.loads(row["section5_json"]) if row["section5_json"] else []
    for i, item in enumerate(s5, start=1):
        student["V_" + str(i)] = item.get("reponse", "")

    # Section B (SEI) : liste
    sb = _json.loads(row["sei_json"]) if row["sei_json"] else []
    for i, item in enumerate(sb, start=1):
        student["B_" + str(i)] = item.get("reponse", "")

    # Totaux
    student["Total_Oui_II"] = row["score_estime_soi"] or 0

    # Calculer Total_Vrai_B : compter les reponses conformes a la polarite
    total_vrai_b = 0
    for item in sb:
        pol = item.get("polarite", "")
        rep = item.get("reponse", "").lower()
        if pol == "vrai_pos" and rep == "vrai":
            total_vrai_b += 1
        elif pol == "vrai_neg" and rep == "faux":
            total_vrai_b += 1
    student["Total_Vrai_B"] = total_vrai_b

    student["PP"] = 0
    student["PN"] = 0
    student["remarques"] = ""

    return render_template(
        "student_view.html", student=student,
        section_ii=SECTION_II, section_iii=SECTION_III, section_iv=SECTION_IV,
        section_v=SECTION_V, section_b=SECTION_B,
    )


@app.route("/students/<int:student_id>/delete", methods=["POST"])
@login_required
@admin_required
def student_delete(student_id):
    conn = db.get_db()
    conn.execute("DELETE FROM students WHERE id = ?", (student_id,))
    conn.commit()
    flash("Fiche supprimée.", "success")
    return redirect(url_for("dashboard"))


# ---------- Export Excel ----------



@app.route("/export.xlsx")
@login_required
def export_xlsx():
    import openpyxl
    import json as _json
    from openpyxl.styles import Font, PatternFill, Alignment

    # === Palette couleurs (identique fichier temoin) ===
    COULEUR_ENTETE      = "1F4E79"
    COULEUR_TRES_ELEVE  = "00B050"
    COULEUR_ELEVE       = "92D050"
    COULEUR_MOYEN       = "FFC000"
    COULEUR_FAIBLE      = "ED7D31"
    COULEUR_TRES_FAIBLE = "FF0000"
    COULEURS_NIVEAU = {
        "Tres faible": COULEUR_TRES_FAIBLE,
        "Faible":      COULEUR_FAIBLE,
        "Moyen":       COULEUR_MOYEN,
        "Eleve":       COULEUR_ELEVE,
        "Tres eleve":  COULEUR_TRES_ELEVE,
    }
    FILL_ENTETE = PatternFill(start_color=COULEUR_ENTETE, end_color=COULEUR_ENTETE, fill_type="solid")
    FONT_ENTETE = Font(color="FFFFFF", bold=True)
    ALIGN_CENTRE = Alignment(horizontal="center", vertical="center", wrap_text=False)

    def _normaliser(v):
        """Uniformise les libelles oui/non/ne_sais_pas."""
        if v is None:
            return ""
        t = str(v).strip()
        bas = t.lower()
        if bas == "oui":
            return "Oui"
        if bas == "non":
            return "Non"
        if bas in ("ne_sais_pas", "ne sais pas", "je ne sais pas"):
            return "Je ne sais pas"
        return t

    def _styler_entete(ws, ligne, nb_cols, hauteur=None):
        for col in range(1, nb_cols + 1):
            c = ws.cell(row=ligne, column=col)
            c.fill = FILL_ENTETE
            c.font = FONT_ENTETE
            c.alignment = ALIGN_CENTRE
        if hauteur:
            ws.row_dimensions[ligne].height = hauteur

    def _colorer_niveau(cellule, valeur):
        if valeur is None:
            return
        v = str(valeur).strip()
        couleur = COULEURS_NIVEAU.get(v)
        if couleur:
            cellule.fill = PatternFill(start_color=couleur, end_color=couleur, fill_type="solid")
            cellule.font = Font(color="FFFFFF", bold=True)
            cellule.alignment = ALIGN_CENTRE

    conn = db.get_db()

    # Recuperer toutes les reponses
    rows = conn.execute(
        "SELECT * FROM reponses ORDER BY LOWER(TRIM(nom)) ASC, LOWER(TRIM(prenom)) ASC"
    ).fetchall()

    # Recuperer les entretiens (pour commentaires SEI)
    ent_rows = conn.execute("SELECT nom_eleve, reponses_json FROM entretiens").fetchall()
    ent_map = {}
    for e in ent_rows:
        nom_e = (e["nom_eleve"] or "").strip().lower()
        try:
            rj = _json.loads(e["reponses_json"]) if e["reponses_json"] else {}
        except:
            rj = {}
        ent_map[nom_e] = rj

    # Preparer les donnees de chaque eleve
    eleves = []
    for r in rows:
        e = dict(r)
        # Section II (40 items oui/non)
        s2 = _json.loads(e["section2_json"]) if e["section2_json"] else []
        for i, item in enumerate(s2, start=1):
            e["II_" + str(i)] = item.get("reponse", "")
        # Section III
        e["III_1"] = e.get("section3_niveau") or ""
        e["III_2"] = e.get("section3_q2") or ""
        e["III_3"] = e.get("section3_q3") or ""
        # Section IV
        s4 = _json.loads(e["section4_json"]) if e["section4_json"] else {}
        for i in range(1, 13):
            e["IV_" + str(i)] = s4.get("q" + str(i), "")
        # Section V
        s5 = _json.loads(e["section5_json"]) if e["section5_json"] else []
        for i, item in enumerate(s5, start=1):
            e["V_" + str(i)] = item.get("reponse", "")
        # Section B (SEI)
        sb = _json.loads(e["sei_json"]) if e["sei_json"] else []
        for i, item in enumerate(sb, start=1):
            e["B_" + str(i)] = item.get("reponse", "")
        # Calculs
        pp2 = sum(1 for a in s2 if a.get("polarite") == "pos" and str(a.get("reponse","")).lower() == "oui")
        pn2 = sum(1 for a in s2 if a.get("polarite") == "neg" and str(a.get("reponse","")).lower() == "non")
        e["PP"] = pp2
        e["PN"] = pn2
        e["Total_Oui_II"] = pp2 + pn2
        tvb = 0
        pp_sei = 0
        pn_sei = 0
        for item in sb:
            pol = item.get("polarite", "")
            rep = str(item.get("reponse","")).lower()
            if pol == "vrai_pos" and rep == "vrai":
                pp_sei += 1
                tvb += 1
            elif pol == "vrai_neg" and rep == "faux":
                pn_sei += 1
                tvb += 1
        e["Total_Vrai_B"] = tvb
        e["PP_SEI"] = pp_sei
        e["PN_SEI"] = pn_sei
        # Niveaux
        s2_score = pp2 + pn2
        if s2_score <= 13: niv2 = "Tres faible"
        elif s2_score <= 20: niv2 = "Faible"
        elif s2_score <= 27: niv2 = "Moyen"
        elif s2_score <= 33: niv2 = "Eleve"
        else: niv2 = "Tres eleve"
        if tvb <= 5: niv_sei = "Tres faible"
        elif tvb <= 8: niv_sei = "Faible"
        elif tvb <= 11: niv_sei = "Moyen"
        elif tvb <= 13: niv_sei = "Eleve"
        else: niv_sei = "Tres eleve"
        e["niv2"] = niv2
        e["niv_sei"] = niv_sei
        e["total"] = s2_score + tvb
        eleves.append(e)

    wb = openpyxl.Workbook()

    # ============================================================
    # FEUILLE 1 : Tableau de bord
    # ============================================================
    ws1 = wb.active
    ws1.title = "Tableau de bord"
    ws1.append(["#", "Nom", "Prenom", "Classe", "Etablissement",
                "PP II", "PN II", "Score II /40", "Niveau II",
                "PP SEI", "PN SEI", "Score SEI /15", "Niveau SEI", "Total /55"])
    for i, e in enumerate(eleves, start=1):
        ws1.append([i, e["nom"], e["prenom"], e["classe"], e["etablissement"],
                    e["PP"], e["PN"], str(e["Total_Oui_II"]) + " / 40", e["niv2"],
                    e["PP_SEI"], e["PN_SEI"], str(e["Total_Vrai_B"]) + " / 15", e["niv_sei"],
                    str(e["total"]) + " / 55"])

    # Stylage en-tete + coloriage des niveaux
    _styler_entete(ws1, 1, 14, hauteur=22)
    for _row in range(2, ws1.max_row + 1):
        _colorer_niveau(ws1.cell(row=_row, column=9), ws1.cell(row=_row, column=9).value)
        _colorer_niveau(ws1.cell(row=_row, column=13), ws1.cell(row=_row, column=13).value)
        _colorer_niveau(ws1.cell(row=_row, column=14), ws1.cell(row=_row, column=14).value)

    # ============================================================
    # FEUILLE 2 : A_Questionnaire
    # ============================================================
    ws2 = wb.create_sheet("A_Questionnaire")
    # Ligne 1 : en-tetes de sections
    h1 = ["I- Informations generales", "", "", "", "", "", "", ""]
    h1 += ["II- Estime de soi de l'eleve"] + [""] * 39
    h1 += ["Totaux", ""]
    h1 += ["III- Perceptions des capacites scolaires", "", ""]
    h1 += ["IV- Orientation scolaire"] + [""] * 11
    h1 += ["V- Influence de l'estime de soi sur l'orientation"] + [""] * 4
    ws2.append(h1)
    # Ligne 2 : noms de colonnes
    h2 = ["ID", "Nom", "Prenoms", "Age", "Sexe", "Classe", "Etablissement", "Situation_familiale"]
    h2 += ["II_" + str(i) for i in range(1, 41)]
    h2 += ["PP", "PN"]
    h2 += ["III_1", "III_2", "III_3"]
    h2 += ["IV_" + str(i) for i in range(1, 13)]
    h2 += ["V_" + str(i) for i in range(1, 6)]
    ws2.append(h2)

    # Stylage lignes 1 et 2
    _styler_entete(ws2, 1, ws2.max_column, hauteur=22)
    _styler_entete(ws2, 2, ws2.max_column, hauteur=30)
    # Donnees
    for e in eleves:
        r = [e["id"], e["nom"], e["prenom"], e["age"], e["sexe"], e["classe"], e["etablissement"], e.get("situation_familiale","")]
        r += [e.get("II_" + str(i), "") for i in range(1, 41)]
        r += [e["PP"], e["PN"]]
        r += [_normaliser(e["III_1"]), _normaliser(e["III_2"]), _normaliser(e["III_3"])]
        r += [_normaliser(e.get("IV_" + str(i), "")) for i in range(1, 13)]
        r += [e.get("V_" + str(i), "") for i in range(1, 6)]
        ws2.append(r)

    # ============================================================
    # FEUILLE 3 : B_Test d'estime de soi
    # ============================================================
    ws3 = wb.create_sheet("B_Test d'estime de soi")
    h1 = ["I- Informations generales", "", "", "", "", "", "", "", ""]
    h1 += ["B- Test d'estime de soi (SEI)"] + [""] * 14
    ws3.append(h1)
    h2 = ["ID", "Nom", "Prenoms", "Age", "Sexe", "Classe", "Etablissement", "Situation_familiale", ""]
    h2 += ["B_" + str(i) for i in range(1, 16)]
    ws3.append(h2)

    _styler_entete(ws3, 1, ws3.max_column, hauteur=22)
    _styler_entete(ws3, 2, ws3.max_column, hauteur=30)
    for e in eleves:
        r = [e["id"], e["nom"], e["prenom"], e["age"], e["sexe"], e["classe"], e["etablissement"], e.get("situation_familiale",""), ""]
        r += [e.get("B_" + str(i), "") for i in range(1, 16)]
        ws3.append(r)

    # ============================================================
    # FEUILLE 4 : C_Guide d'entretien
    # ============================================================
    ws4 = wb.create_sheet("C_Guide d'entretien")
    h1 = ["I- Informations generales", "", "", "", "", "", "", "", ""]
    h1 += ["C- Guide d'entretien"] + [""] * 11
    ws4.append(h1)
    h2 = ["ID", "Nom", "Prenoms", "Age", "Sexe", "Classe", "Etablissement", "Situation_familiale", ""]
    h2 += ["1-", "2-a", "2-b", "2-c", "3-a", "3-b", "4-a", "4-b", "5", "6", "7", "8"]
    ws4.append(h2)

    _styler_entete(ws4, 1, ws4.max_column, hauteur=22)
    _styler_entete(ws4, 2, ws4.max_column, hauteur=30)
    # Colonnes de correspondance entre les cles JSON des entretiens et les colonnes Excel
    guide_keys = ["q1", "q2a", "q2b", "q2c", "q3a", "q3b", "q4a", "q4b", "q5", "q6", "q7", "q8"]
    for e in eleves:
        r = [e["id"], e["nom"], e["prenom"], e["age"], e["sexe"], e["classe"], e["etablissement"], e.get("situation_familiale",""), ""]
        # Chercher l'entretien correspondant
        nom_complet = (str(e["nom"]) + " " + str(e["prenom"])).strip().lower()
        rj = ent_map.get(nom_complet, {})
        for key in guide_keys:
            r.append(rj.get(key, ""))
        ws4.append(r)

    # Sauvegarder
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    filename = f"Depouillement_Enquete_EstimeDeSoi_{datetime.now():%Y%m%d_%H%M%S}.xlsx"
    return send_file(
        buf, as_attachment=True, download_name=filename,
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )



@app.route("/admin/import", methods=["GET", "POST"])
@login_required
@admin_required
def admin_import():
    if request.method == "POST":
        file = request.files.get("xlsx_file")
        if not file or file.filename == "":
            flash("Aucun fichier sélectionné.", "error")
            return redirect(url_for("admin_import"))

        import openpyxl
        wb = openpyxl.load_workbook(file, data_only=True)
        if "Données" not in wb.sheetnames:
            flash('Feuille "Données" introuvable dans ce fichier.', "error")
            return redirect(url_for("admin_import"))

        ws = wb["Données"]
        headers = [c.value for c in ws[1]]
        conn = db.get_db()
        imported = 0

        for row in ws.iter_rows(min_row=2, values_only=True):
            row_dict = dict(zip(headers, row))
            # Ignore les lignes vides (pas de nom)
            if not row_dict.get("Nom"):
                continue

            db_row = {}
            for col_name, value in row_dict.items():
                if col_name in (None, "ID"):
                    continue
                db_field = _IMPORT_FIELD_MAP.get(col_name, col_name)
                db_row[db_field] = value

            db_row["created_by"] = current_user.id
            columns_sql = list(db_row.keys())
            placeholders = ", ".join(["?"] * len(columns_sql))
            conn.execute(
                f"INSERT INTO students ({', '.join(columns_sql)}) VALUES ({placeholders})",
                list(db_row.values()),
            )
            imported += 1

        conn.commit()
        flash(f"{imported} fiche(s) importée(s) avec succès.", "success")
        return redirect(url_for("dashboard"))

    return render_template("admin_import.html")


# ---------- Gestion des utilisateurs (admin) ----------

@app.route("/admin/users", methods=["GET", "POST"])
@login_required
@admin_required
def admin_users():
    conn = db.get_db()
    if request.method == "POST":
        username = request.form["username"].strip()
        password = request.form["password"]
        role = request.form.get("role", "enqueteur")
        try:
            conn.execute(
                "INSERT INTO users (username, password_hash, role) VALUES (?, ?, ?)",
                (username, generate_password_hash(password), role),
            )
            conn.commit()
            flash(f"Utilisateur {username} créé.", "success")
        except Exception as e:
            flash(f"Erreur : {e}", "error")
    users = conn.execute("SELECT id, username, role, created_at FROM users").fetchall()
    return render_template("admin_users.html", users=users)




# ==================== ENTRE-TIENS ====================

def can_rename_required(fn):
    from functools import wraps
    @wraps(fn)
    def wrapper(*args, **kwargs):
        if not current_user.is_authenticated or current_user.role not in ("admin", "editeur"):
            abort(403)
        return fn(*args, **kwargs)
    return wrapper


@app.route("/entretien", methods=["GET", "POST"])
@login_required
def entretien():
    import json as _json
    conn = db.get_db()
    from items import GUIDE_ENTRETIEN, SEI_COMMENTAIRES

    # Recuperer les eleves pour le selecteur
    eleves = conn.execute("SELECT id, nom, prenom FROM reponses ORDER BY LOWER(TRIM(nom)), LOWER(TRIM(prenom))").fetchall()

    # Si un eleve est selectionne (pour afficher ses reponses SEI)
    eleve_id = request.args.get("eleve_id", type=int)
    sei_reponses = {}
    if eleve_id:
        row = conn.execute("SELECT sei_json FROM reponses WHERE id=?", (eleve_id,)).fetchone()
        if row and row["sei_json"]:
            try:
                sei_data = _json.loads(row["sei_json"])
                for i, item in enumerate(sei_data):
                    rep = item.get("reponse", "")
                    pol = item.get("polarite", "")
                    # Determiner la "valeur" attendue
                    sei_reponses["sei_comment_" + str(i)] = rep
            except:
                pass

    if request.method == "POST":
        nom_eleve = request.form.get("nom_eleve", "").strip()
        reponses = {}
        for key, _ in GUIDE_ENTRETIEN:
            reponses[key] = request.form.get(key, "").strip()
        for key, _ in SEI_COMMENTAIRES:
            reponses[key] = request.form.get(key, "").strip()
        if nom_eleve:
            conn.execute(
                "INSERT INTO entretiens (date_entretien, nom_eleve, reponses_json) VALUES (datetime('now'), ?, ?)",
                (nom_eleve, _json.dumps(reponses, ensure_ascii=False)),
            )
            conn.commit()
            flash("Entretien enregistre.", "success")
            return redirect(url_for("entretien"))

    return render_template(
        "entretien.html",
        questions=GUIDE_ENTRETIEN,
        sei_commentaires=SEI_COMMENTAIRES,
        sei_reponses=sei_reponses,
        eleves=eleves,
    )


@app.route("/admin/entretiens")
@login_required
def admin_entretiens():
    conn = db.get_db()
    entretiens = conn.execute(
        "SELECT id, date_entretien, nom_eleve FROM entretiens ORDER BY LOWER(TRIM(nom_eleve)) ASC, id DESC"
    ).fetchall()
    return render_template("admin_entretiens.html", entretiens=entretiens)


@app.route("/admin/renommer_entretien/<int:ent_id>", methods=["POST"])
@login_required
@can_rename_required
def renommer_entretien(ent_id):
    nom = request.form.get("nom_eleve", "").strip()
    if nom:
        conn = db.get_db()
        conn.execute("UPDATE entretiens SET nom_eleve=? WHERE id=?", (nom, ent_id))
        conn.commit()
    return redirect(url_for("admin_entretiens"))


@app.route("/admin/supprimer_entretien/<int:ent_id>", methods=["POST"])
@login_required
@admin_required
def supprimer_entretien(ent_id):
    conn = db.get_db()
    conn.execute("DELETE FROM entretiens WHERE id=?", (ent_id,))
    conn.commit()
    flash("Entretien supprime.", "success")
    return redirect(url_for("admin_entretiens"))


@app.route("/entretien/modifier/<int:ent_id>", methods=["GET", "POST"])
@login_required
@admin_required
def modifier_entretien(ent_id):
    import json as _json
    conn = db.get_db()
    from items import GUIDE_ENTRETIEN, SEI_COMMENTAIRES
    ent = conn.execute("SELECT * FROM entretiens WHERE id=?", (ent_id,)).fetchone()
    if not ent:
        abort(404)
    if request.method == "POST":
        nom_eleve = request.form.get("nom_eleve", "").strip()
        reponses = {}
        for key, _ in GUIDE_ENTRETIEN:
            reponses[key] = request.form.get(key, "").strip()
        for key, _ in SEI_COMMENTAIRES:
            reponses[key] = request.form.get(key, "").strip()
        conn.execute(
            "UPDATE entretiens SET nom_eleve=?, reponses_json=? WHERE id=?",
            (nom_eleve, _json.dumps(reponses, ensure_ascii=False), ent_id),
        )
        conn.commit()
        flash("Entretien modifie.", "success")
        return redirect(url_for("admin_entretiens"))
    reponses = _json.loads(ent["reponses_json"]) if ent["reponses_json"] else {}

    # Charger les reponses SEI de l'eleve correspondant (par nom)
    sei_reponses = {}
    if ent["nom_eleve"]:
        # Rechercher l'eleve par nom complet
        parts = ent["nom_eleve"].strip().split(" ", 1)
        if len(parts) == 2:
            nom, prenom = parts[0], parts[1]
            row = conn.execute(
                "SELECT sei_json FROM reponses WHERE LOWER(TRIM(nom))=LOWER(TRIM(?)) AND LOWER(TRIM(prenom))=LOWER(TRIM(?)) LIMIT 1",
                (nom, prenom)
            ).fetchone()
            if row and row["sei_json"]:
                try:
                    sei_data = _json.loads(row["sei_json"])
                    for i, item in enumerate(sei_data):
                        sei_reponses["sei_comment_" + str(i)] = item.get("reponse", "")
                except:
                    pass

    return render_template(
        "entretien.html",
        questions=GUIDE_ENTRETIEN,
        sei_commentaires=SEI_COMMENTAIRES,
        sei_reponses=sei_reponses,
        ent=ent,
        reponses=reponses,
        modifier=True,
    )


@app.cli.command("create-admin")
def create_admin_command():
    """Commande CLI pour créer le premier compte admin : flask create-admin"""
    import getpass
    username = input("Nom d'utilisateur admin : ").strip()
    password = getpass.getpass("Mot de passe : ")
    with app.app_context():
        conn = db.get_db()
        conn.execute(
            "INSERT INTO users (username, password_hash, role) VALUES (?, ?, 'admin')",
            (username, generate_password_hash(password)),
        )
        conn.commit()
    print(f"Admin {username} créé.")


with app.app_context():
    db.init_db(app)
    # Crée un compte admin par défaut si la base est vierge (à changer immédiatement)
    conn = db.get_db()
    count = conn.execute("SELECT COUNT(*) AS c FROM users").fetchone()["c"]
    if count == 0:
        default_user = os.environ.get("ADMIN_USERNAME", "admin")
        default_pass = os.environ.get("ADMIN_PASSWORD", "admin123")
        conn.execute(
            "INSERT INTO users (username, password_hash, role) VALUES (?, ?, 'admin')",
            (default_user, generate_password_hash(default_pass)),
        )
        conn.commit()


if __name__ == "__main__":
    app.run(host="0.0.0.0", port=5000, debug=False)


# ---------- Suppression / modification des fiches eleves ----------

@app.route("/admin/supprimer/<int:eleve_id>", methods=["POST"])
@login_required
@admin_required
def supprimer_eleve(eleve_id):
    conn = db.get_db()
    conn.execute("DELETE FROM reponses WHERE id=?", (eleve_id,))
    conn.commit()
    flash("Fiche supprimee.", "success")
    return redirect(url_for("dashboard"))


@app.route("/admin/supprimer_selection", methods=["POST"])
@login_required
@admin_required
def supprimer_selection():
    conn = db.get_db()
    ids_rep = request.form.getlist("selection_reponses")
    ids_ent = request.form.getlist("selection_entretiens")
    for rid in ids_rep:
        conn.execute("DELETE FROM reponses WHERE id=?", (rid,))
    for eid in ids_ent:
        conn.execute("DELETE FROM entretiens WHERE id=?", (eid,))
    conn.commit()
    flash(f"{len(ids_rep) + len(ids_ent)} element(s) supprime(s).", "success")
    return redirect(url_for("dashboard"))


@app.route("/admin/modifier/<int:eleve_id>", methods=["GET", "POST"])
@login_required
@admin_required
def modifier_eleve(eleve_id):
    import json as _json
    conn = db.get_db()
    row = conn.execute("SELECT * FROM reponses WHERE id=?", (eleve_id,)).fetchone()
    if not row:
        abort(404)
    if request.method == "POST":
        # Mise a jour des champs simples
        conn.execute(
            "UPDATE reponses SET nom=?, prenom=?, age=?, sexe=?, classe=?, etablissement=?, situation_familiale=? WHERE id=?",
            (
                request.form.get("nom"),
                request.form.get("prenom"),
                request.form.get("age"),
                request.form.get("sexe"),
                request.form.get("classe"),
                request.form.get("etablissement"),
                request.form.get("situation_familiale"),
                eleve_id,
            ),
        )
        conn.commit()
        flash("Fiche modifiee.", "success")
        return redirect(url_for("dashboard"))

    # GET : afficher le formulaire (reutilise student_form.html)
    student = dict(row)
    return render_template(
        "student_form.html",
        student=student,
        section_ii=SECTION_II, section_iii=SECTION_III, section_iv=SECTION_IV,
        section_v=SECTION_V, section_b=SECTION_B,
        age_choices=AGE_CHOICES, sexe_choices=SEXE_CHOICES,
        situation_choices=SITUATION_FAMILIALE_CHOICES,
    )


# ---------- Gestion des erreurs ----------

@app.errorhandler(403)
def forbidden(e):
    return render_template("403.html"), 403


@app.errorhandler(404)
def not_found(e):
    return render_template("403.html"), 404


# ---------- Capture d'erreurs ----------
import traceback
import sys

@app.errorhandler(500)
def internal_error(e):
    with open('/app/data/errors.log', 'a') as f:
        f.write('=== ' + str(datetime.now()) + ' ===\n')
        f.write(str(e) + '\n')
        if hasattr(e, 'original_exception'):
            traceback.print_exception(type(e.original_exception), e.original_exception, e.original_exception.__traceback__, file=f)
        f.write('\n')
    return "Erreur interne - voir /app/data/errors.log", 500
